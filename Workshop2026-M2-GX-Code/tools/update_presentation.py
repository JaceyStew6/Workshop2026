"""Met à jour la présentation (Workshop2026-M2-GX-Pres.pptx) avec les résultats mesurés sous Docker.

  python tools/update_presentation.py --src ancienne.pptx --out nouvelle.pptx

Les textes sont modifiés en place (le style de chaque zone est conservé) ; le graphique natif de la diapo 5
reçoit les nouvelles données, son image est remplacée par tests/rehearsals/latency-run-1.png.
Dépendances : pip install python-pptx pillow.
"""
import argparse
import copy
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"{name} introuvable sur la diapositive")


def set_paragraphs(sh, texts):
    """Remplace le texte paragraphe par paragraphe en gardant la mise en forme du premier segment de chacun."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    for i, text in enumerate(texts):
        if i < len(paras):
            para = paras[i]
        else:  # copie du dernier paragraphe pour hériter de son style
            new = copy.deepcopy(paras[-1]._p)
            paras[-1]._p.addnext(new)
            paras = list(tf.paragraphs)
            para = paras[i]
        runs = para.runs
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    for extra in paras[len(texts):]:
        extra._p.getparent().remove(extra._p)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    prs = Presentation(a.src)
    S = {i + 1: s for i, s in enumerate(prs.slides)}

    runs = [json.loads((ROOT / "tests" / "rehearsals" / f"run-{n}" / "result.json").read_text(encoding="utf-8")) for n in (1, 2, 3)]
    det = [r["incident"]["detected_after_s"] for r in runs]
    rep = [r["incident"]["detect_to_recovered_s"] for r in runs]
    rec = [r["incident"]["recovered_after_s"] for r in runs]
    assert all(r["verdict"] == "REUSSI" for r in runs), "une répétition a échoué : ne pas annoncer 3/3"

    # Diapo 2 : le scénario
    set_paragraphs(shape(S[2], "Text 6"), ["≈ 20 ms → 2 s"])
    set_paragraphs(shape(S[2], "Text 7"), ["latence perçue par un client, avec erreurs 503"])

    # Diapo 4 : la chaîne
    set_paragraphs(shape(S[4], "Text 4"), ["Isolation Forest sur 7 mesures ; 3 échantillons sur 4"])
    set_paragraphs(shape(S[4], "Text 16"), ["Transaction réelle, aucun argent perdu, 8 s de stabilité"])

    # Diapo 6 : le graphique vient du banc natif antérieur, le dire sur la diapositive
    set_paragraphs(shape(S[6], "Text 3"), ["Source impossible à bannir",
                                           "Un redémarrage donne quelques secondes de répit, puis l'attaque reprend.",
                                           "Résultat : escalade humaine après 49 s (banc natif)."])

    # Diapo 5 : résultats mesurés (Docker)
    set_paragraphs(shape(S[5], "Message"), ["Trois répétitions sous Docker, sans intervention : l'attaque continue après la reprise."])
    set_paragraphs(shape(S[5], "Text 2"), [f"≈ {round(sum(det) / len(det))} s"])
    set_paragraphs(shape(S[5], "Text 4"), [f"≈ {round(sum(rep) / len(rep))} s"])
    set_paragraphs(shape(S[5], "Text 8"), [
        "Docker, une seule machine, attaquant simulé : valide la chaîne, pas la capture de l'IP d'un second PC. "
        "Exemple : répétition 1 (300 req/s)."])
    chart = shape(S[5], "Chart 0").chart
    cd = CategoryChartData()
    cd.categories = ["Essai 1", "Essai 2", "Essai 3"]
    cd.add_series("Détection", det)
    cd.add_series("Reprise", rep)
    chart.replace_data(cd)
    pic = shape(S[5], "Image 0")
    img = ROOT / "tests" / "rehearsals" / "latency-run-1.png"
    part = S[5].part.related_part(pic._element.blipFill.blip.rEmbed)
    part._blob = img.read_bytes()
    w, h = Image.open(img).size
    pic.height = Emu(int(pic.width * h / w))
    set_notes(S[5],
        f"Chiffres (3 répétitions sous Docker, attaquant simulé sur la même machine) : détection {det[0]} / {det[1]} / {det[2]} s ; "
        f"reprise {rec[0]} / {rec[1]} / {rec[2]} s après le début de l'attaque, soit {rep[0]} / {rep[1]} / {rep[2]} s après la détection. "
        "L'attaque dure 120 s : elle continue plus de 90 s après la reprise, donc c'est bien la remédiation qui agit, pas la fin du script. "
        "Pendant l'incident, devis et virements échouent en 503 pour tous les clients ; ensuite 100 % des virements du client légitime réussissent. "
        "Dire clairement que l'attaquant est simulé (en-tête X-Forwarded-For) et que l'essai depuis un second PC reste à faire.")

    # Diapo 6 : attaquant caché (notes seulement : l'image vient du banc natif)
    set_notes(S[6],
        "Moment d'honnêteté. Cette image vient du banc natif antérieur. Le bon comportement face à une source impossible à bannir est d'escalader. "
        "Constaté ensuite sous Docker Desktop : sur un seul poste, le WAF ne voit que la passerelle du réseau Docker pour tout le trafic de l'hôte, "
        "donc un attaquant lancé depuis le même PC est indiscernable d'un client ; en laboratoire nous simulons des identités distinctes, "
        "mais l'IP réelle d'un second PC reste à vérifier. Montrer que les essais ont servi à corriger de vrais défauts.")

    # Diapo 7 : IA et données
    set_paragraphs(shape(S[7], "Text 5"), ["Isolation Forest, 7 mesures Prometheus", "Seuil au quantile 99,9 % de la validation",
                                           "Alerte seulement si 3 échantillons sur 4", "Aucune alerte si une métrique manque"])
    set_paragraphs(shape(S[7], "Text 10"), ["Limite assumée : 132 min de trafic réel (actif et inactif) ; à réentraîner après tout changement de l'hôte."])
    set_notes(S[7],
        "Découvertes avec les données réelles : les séries d'erreurs sont absentes quand il n'y a pas d'erreur (détecteur aveugle, corrigé) ; "
        "l'absence de trafic doit faire partie du normal ; une variable quasi constante (la mémoire) faussait tout le seuil et a été retirée ; "
        "un faux incident pendant la collecte a même banni notre propre générateur de trafic, d'où le filtre anti-contamination. "
        "Au repos le score reste entre 0,40 et 0,58 pour un seuil de 0,70 ; en attaque il atteint 0,77. "
        "Le RAG a été testé avec un faux modèle (10 tests) ; Ollama n'a pas été exécuté.")

    # Diapo 8 : cybersécurité
    set_paragraphs(shape(S[8], "Text 13"), ["6 règles d'alerte MITRE ATT&CK validées par promtool ; journaux JSON dans Loki (vérifié en direct)."])

    # Diapo 9 : gouvernance
    set_paragraphs(shape(S[9], "Text 16"), [
        "PCA/PRA : 6 scénarios, ordre de reconstruction, sauvegarde et restauration rejouées sous Docker (Ansible, 19 s). "
        "NIS2 et DORA : cartographie couvert / partiel / écart, sans prétention de conformité."])

    # Diapo 10 : ce qui reste à prouver (le point 2, en rouge, reste le risque principal)
    items = [("Text 4", "Text 5", "Assistant RAG avec Ollama réel", "Installer Ollama, calibrer le seuil de pertinence"),
             ("Text 8", "Text 9", "Attaque depuis un second PC", "Lire l'IP source vue par le WAF, TRUSTED_PROXIES vide"),
             ("Text 12", "Text 13", "Détecteur sur le dispositif final", "Réentraîner (≥ 2 h), puis répéter 3 fois depuis le PC Attaquant"),
             ("Text 16", "Text 17", "Vidéo de 60 s sans coupe", "Filmer une exécution réelle depuis le PC Attaquant"),
             ("Text 20", "Text 21", "TLS, MFA, scan d'images", "Écarts assumés, premiers sujets d'une feuille de route")]
    for left, right, title, action in items:
        set_paragraphs(shape(S[10], left), [title])
        set_paragraphs(shape(S[10], right), [action])
    set_paragraphs(shape(S[10], "Message"), ["Ce qui reste à prouver sur le dispositif final, et comment nous le vérifions."])
    set_notes(S[10],
        "Poser le cadre du jury : nous préférons dire ce qui n'est pas prouvé. Prouvé depuis la dernière version : la composition Docker (16 services), "
        "le déploiement Ansible avec secrets chiffrés, Loki, le détecteur entraîné sur 132 minutes et trois répétitions réussies. "
        "Le point 2 (second PC, IP source) conditionne tout le bannissement par IP : sur un seul poste le WAF ne voit que la passerelle.")

    prs.save(a.out)
    print("écrit :", a.out)


if __name__ == "__main__":
    main()
